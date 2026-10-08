import re, sys, datetime as dt
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList

OUT = sys.argv[1]
TRACKER = sys.argv[2] if len(sys.argv) > 2 else "Sales_Tracker_TSX.xlsx"
ROSTER = sys.argv[3] if len(sys.argv) > 3 else "TSX_Updated_Roster.xlsx"

# ----------------------------------------------------------------- load & clean
s = pd.read_excel(TRACKER)
r = pd.read_excel(ROSTER, sheet_name="Roster")
r.columns = [str(c).strip() for c in r.columns]
r = r[r["LOB"].notna()].copy()
r["HRID"] = r["HRID"].astype(str).str.extract(r"(\d+)")[0]
r["Agent Name"] = r["Agent Name"].astype(str).str.strip().str.title()
r["Team Leader"] = r["Team Leader"].astype(str).str.strip()
r["Ops Manager"] = r["Ops Manager"].astype(str).str.strip()
r["Tenure Days"] = pd.to_numeric(r["Tenurity"], errors="coerce")
r["LOB Start Date"] = pd.to_datetime(r["LOB Start Date"], errors="coerce")
r = r.sort_values(["Team Leader", "Agent Name"]).reset_index(drop=True)
r["Label"] = r["Agent Name"] + " (" + r["HRID"] + ")"
roster = set(r["HRID"])
rmap = r.set_index("HRID")

def map_hrid(x):
    x = re.sub(r"\D", "", str(x).strip())
    if x in roster: return x
    if len(x) == 8 and x.startswith("100") and ("1000" + x[3:]) in roster: return "1000" + x[3:]
    if len(x) == 10 and x.startswith("10000") and ("1000" + x[5:]) in roster: return "1000" + x[5:]
    if len(x) == 10 and x.startswith("1000") and x[:9] in roster: return x[:9]
    return None

s["ct"] = pd.to_datetime(s["Completion time"])
s["Date"] = s["ct"].dt.normalize()
s["hr_raw"] = s["HRID"].fillna("").astype(str).str.strip()
s["hr"] = s["hr_raw"].apply(map_hrid)
s["in_roster"] = s["hr"].notna().astype(int)
s["hr_key"] = s["hr"].fillna("UNMAPPED:" + s["hr_raw"].str.replace(r"\s+", " ", regex=True))
s["Agent"] = s["hr"].map(rmap["Agent Name"]).fillna("Not in Roster")
s["TL"] = s["hr"].map(rmap["Team Leader"]).fillna("Not in Roster")
s["OM"] = s["hr"].map(rmap["Ops Manager"]).fillna("Not in Roster")
s["Phase"] = s["hr"].map(rmap["Phase"]).fillna("-")
s["Tenure"] = s["hr"].map(rmap["Tenure Bucket"]).fillna("-")
s["Type"] = s["Line or GSR"].str.strip()
s["xm"] = np.where(s["Type"] == "XM Lines", s["No. of XM Lines"].fillna(1), 0).astype(int)
types = s["Type of GSR"].fillna("").str.strip().str.strip(";")
ntypes = types.apply(lambda t: len([p for p in t.split(";") if p.strip()]))
gcount = s["Count of GSR"]
s["gsr"] = np.where(s["Type"] == "GSRs", np.where(gcount.notna(), gcount.fillna(0), np.maximum(ntypes, 1)), 0).astype(int)
s["units"] = s["xm"] + s["gsr"]
def typeflag(name):
    has = types.str.contains(name, regex=False)
    return np.where(has, np.where(ntypes == 1, s["gsr"], 1), 0).astype(int)
s["t_speed"] = typeflag("Speed Upgrade"); s["t_xpro"] = typeflag("X-PRO")
s["t_lob"] = typeflag("Adding LOB"); s["t_mplus"] = typeflag("Mobile Plus"); s["t_xmc"] = typeflag("XMC")
def clean_dev(v):
    if pd.isna(v): return ""
    t = str(v).strip().lower().replace("\xa0", " ")
    if "iphone" in t: return "iPhone"
    if "apple watch" in t: return "Apple Watch"
    if "samsung watch" in t: return "Samsung Watch"
    if "samsung" in t: return "Samsung"
    if "ipad" in t: return "iPad"
    if "motorola" in t: return "Motorola"
    if "byod" in t: return "BYOD"
    if t in {".", "no", "none", "0", "na", "1", "no ", "n/a"} or "no device" in t or "not solved" in t or "own device" in t or "no devices" in t or "just the line" in t: return "No Device"
    return "Other"
s["device"] = s["Sold Device Type"].apply(clean_dev)
def clean_promo(v):
    if pd.isna(v): return ""
    t = str(v).strip().lower()
    if "free line" in t: return "Free Line for 1 Year"
    if "byod" in t: return "BYOD"
    if "smart watch" in t or "watch" in t: return "Smart Watch Discount"
    if "port" in t: return "Port In"
    if "no promotion" in t: return "New Line with no promotions"
    if "trade" in t: return "Trade In"
    return "Other"
s["promo"] = s["What was the promotion type"].apply(clean_promo)
s["acct"] = s["Account Number:"].fillna("").astype(str).str.strip()
s["order"] = s["Order Number"].fillna("").astype(str).str.strip()
s["order_valid"] = s["order"].str.fullmatch(r"\d{6,20}").fillna(False).astype(int)
s = s.sort_values("ct").reset_index(drop=True)
key = s["hr_key"] + "|" + s["acct"] + "|" + s["order"] + "|" + s["Type"]
s["dup"] = ((s["order_valid"] == 1) & key.duplicated(keep="first")).astype(int)
nd = s[s["dup"] == 0]
first = ~nd.duplicated(subset=["Date", "hr_key"], keep="first")
s["day_first"] = 0; s.loc[nd.index[first], "day_first"] = 1
s["Weekday"] = s["Date"].dt.strftime("%a")
for c in ["Did you add XMC?", "What is the type of the SIM", "Did you Activate the E-SIM?"]:
    s[c] = s[c].fillna("").astype(str).str.strip()

NS = len(s); LAST = NS + 1
dates = sorted(s["Date"].dt.date.unique()); ND = len(dates)
tls = sorted(r["Team Leader"].unique()); oms = sorted(r["Ops Manager"].unique())
MAXD = max(dates); DEF_TO = MAXD; DEF_FROM = MAXD - dt.timedelta(days=6)
print("rows", NS, "dates", ND, "tls", len(tls), "unmatched", (s.in_roster == 0).sum(), "dups", s.dup.sum())

# ----------------------------------------------------------------- styles
NAVY = "1F3864"; TEAL = "2E75B6"; LIGHT = "D9E1F2"; GREY = "F2F2F2"; INPUT = "FFF2CC"; GREEN = "1E8449"; RED = "C0392B"
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
def F(size=11, bold=False, color="000000", italic=False): return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
def put(ws, ref, v, font=None, fl=None, fmt=None, border=True):
    c = ws[ref]; c.value = v
    c.font = font or F()
    if fl: c.fill = fill(fl)
    if fmt: c.number_format = fmt
    if border: c.border = BORDER
    return c
def header_row(ws, row, c1, labels, fl=NAVY, color="FFFFFF", size=11):
    for i, lab in enumerate(labels):
        put(ws, f"{L(c1+i)}{row}", lab, F(size, True, color), fl).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
def title(ws, text, sub=None, width=14):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    put(ws, "A1", text, F(18, True, NAVY), border=False)
    ws.row_dimensions[1].height = 30
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        put(ws, "A2", sub, F(10, False, "595959", True), border=False)
def section(ws, row, text, c1=1, c2=14):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    put(ws, f"{L(c1)}{row}", text, F(12, True, "FFFFFF"), TEAL)
    ws.row_dimensions[row].height = 20
def widths(ws, d):
    for k, v in d.items(): ws.column_dimensions[k].width = v

# ----------------------------------------------------------------- ranges
def rg(col): return f"Sales_Data!${col}${2}:${col}${LAST}"
C_DATE, C_HR, C_TL, C_OM, C_XM, C_GSR, C_U, C_XMC, C_ESIM, C_DEV, C_DUP, C_INR, C_INC, C_DF = (rg(c) for c in ["C","F","H","I","M","N","O","U","W","X","AC","AD","AE","AF"])
D_FROM, D_TO, D_OM, D_TL, D_DUPT, D_DAYS, D_PFROM, D_PTO, D_TLC, D_OMC = ("Dashboard!$C$5","Dashboard!$C$6","Dashboard!$C$7","Dashboard!$C$8","Dashboard!$C$9","Dashboard!$C$10","Dashboard!$C$12","Dashboard!$C$13","Dashboard!$C$14","Dashboard!$C$15")
RN = len(r); RL = RN + 1
R_HR, R_NAME, R_LABEL, R_TL, R_OM, R_PHASE, R_KEY = (f"Roster!${c}$2:${c}${RL}" for c in ["A","B","C","D","E","F","K"])
AS1, ASL = 5, 4 + RN
def asr(col): return f"Agent_Summary!${col}${AS1}:${col}${ASL}"
METRICS = ["Total Units", "XM Lines", "GSR Units", "Submissions"]
def metric_choose(sel, extra_crit):
    # extra_crit: list of (range, crit) pairs as strings 'range,crit'
    ec = "".join(f",{a},{b}" for a, b in extra_crit)
    return (f"CHOOSE(MATCH({sel},Lists!$C$2:$C$5,0),"
            f"SUMIFS({C_U}{ec},{C_INC},1),SUMIFS({C_XM}{ec},{C_INC},1),SUMIFS({C_GSR}{ec},{C_INC},1),SUMIFS({C_INC}{ec}))")

wb = Workbook()
# ================================================================= Sales_Data
sd = wb.active; sd.title = "Sales_Data"
cols = ["ID","Completion Time","Date","Weekday","HRID (Raw)","HRID (Mapped)","Agent Name","Team Leader","Ops Manager","Phase","Tenure Bucket","Entry Type","XM Lines","GSR Units","Total Units","Speed Upgrade","X-PRO","Adding LOB","Mobile Plus Plan","XMC (GSR)","Added XMC?","SIM Type","E-SIM Activated?","Device Sold","Promotion","Account Number","Order Number","Order Valid","Duplicate Flag","In Roster","Include","Agent-Day First"]
header_row(sd, 1, 1, cols)
rows = []
for i, x in s.iterrows():
    rows.append([int(x["ID"]), x["ct"].to_pydatetime(), x["Date"].date(), x["Weekday"], x["hr_raw"], x["hr_key"], x["Agent"], x["TL"], x["OM"], x["Phase"], x["Tenure"], x["Type"], int(x["xm"]), int(x["gsr"]), int(x["units"]), int(x["t_speed"]), int(x["t_xpro"]), int(x["t_lob"]), int(x["t_mplus"]), int(x["t_xmc"]), x["Did you add XMC?"], x["What is the type of the SIM"], x["Did you Activate the E-SIM?"], x["device"], x["promo"], x["acct"], x["order"][:60], int(x["order_valid"]), int(x["dup"]), int(x["in_roster"]), None, int(x["day_first"])])
for ri, row in enumerate(rows, start=2):
    for ci, v in enumerate(row, start=1):
        c = sd.cell(row=ri, column=ci, value=v); c.font = F(10)
    sd.cell(row=ri, column=2).number_format = "dd-mmm-yy hh:mm"
    sd.cell(row=ri, column=3).number_format = "dd-mmm-yy"
    sd.cell(row=ri, column=31, value=f'=IF(AND({D_DUPT}="Yes",AC{ri}=1),0,1)').font = F(10)
sd.freeze_panes = "A2"; sd.auto_filter.ref = f"A1:AF{LAST}"
widths(sd, {"A":7,"B":17,"C":11,"D":8,"E":14,"F":16,"G":28,"H":17,"I":16,"J":11,"K":11,"L":10,"M":9,"N":9,"O":10,"P":9,"Q":8,"R":9,"S":10,"T":9,"U":10,"V":12,"W":12,"X":13,"Y":22,"Z":18,"AA":22,"AB":9,"AC":9,"AD":9,"AE":8,"AF":9})
sd.sheet_properties.tabColor = "7F7F7F"

# ================================================================= Roster
ro = wb.create_sheet("Roster")
header_row(ro, 1, 1, ["HRID","Agent Name","Agent Label","Team Leader","Ops Manager","Phase","Wave","Tenure Bucket","LOB Start Date","Tenure Days","Team Key"])
for i, x in r.iterrows():
    ri = i + 2
    vals = [x["HRID"], x["Agent Name"], x["Label"], x["Team Leader"], x["Ops Manager"], x["Phase"], x["Wave"], x["Tenure Bucket"], (x["LOB Start Date"].date() if pd.notna(x["LOB Start Date"]) else "-"), (int(x["Tenure Days"]) if pd.notna(x["Tenure Days"]) else "-"), f'=D{ri}&"|"&COUNTIF($D$2:D{ri},D{ri})']
    for ci, v in enumerate(vals, start=1):
        put(ro, f"{L(ci)}{ri}", v, F(10))
    ro[f"I{ri}"].number_format = "dd-mmm-yy"
ro.freeze_panes = "A2"; ro.auto_filter.ref = f"A1:K{RL}"
widths(ro, {"A":12,"B":34,"C":40,"D":18,"E":16,"F":11,"G":13,"H":12,"I":13,"J":11,"K":22})
ro.sheet_properties.tabColor = "7F7F7F"

# ================================================================= Lists
li = wb.create_sheet("Lists")
header_row(li, 1, 1, ["Team Leaders","Ops Managers","Metrics","Agents","Yes/No"])
for i, v in enumerate(["All"] + tls, start=2): put(li, f"A{i}", v, F(10))
for i, v in enumerate(["All"] + oms, start=2): put(li, f"B{i}", v, F(10))
for i, v in enumerate(METRICS, start=2): put(li, f"C{i}", v, F(10))
labels_sorted = sorted(r["Label"].tolist())
for i, v in enumerate(labels_sorted, start=2): put(li, f"D{i}", v, F(10))
for i, v in enumerate(["Yes","No"], start=2): put(li, f"E{i}", v, F(10))
widths(li, {"A":20,"B":18,"C":14,"D":42,"E":8}); li.sheet_properties.tabColor = "7F7F7F"

# ================================================================= Daily_Account
da = wb.create_sheet("Daily_Account")
title(da, "Account Level – Day-over-Day (Cairo Repair TSX)", "Whole-account view, every day in the tracker. Duplicate handling follows the Dashboard toggle. 7-Day Avg and WoW % are trailing calculations.", 26)
dcols = ["Date","Weekday","Submissions","XM Lines","GSR Units","Total Units","DoD Δ Units","DoD %","7-Day Avg","WoW %","Active Sellers","Units / Seller","XMC Yes","XMC Asked","XMC Attach %","E-SIM Activated","E-SIM Asked","E-SIM Act. %","Speed Upgrade","X-PRO","Adding LOB","Mobile Plus","XMC (GSR)","Devices Sold","Not-in-Roster Entries","Duplicates Flagged"]
header_row(da, 3, 1, dcols); da.row_dimensions[3].height = 32
DA1 = 4; DAL = DA1 + ND - 1
for k, d in enumerate(dates):
    rr = DA1 + k
    put(da, f"A{rr}", d, F(10, True), fmt="dd-mmm-yy"); put(da, f"B{rr}", d.strftime("%a"), F(10))
    f = {
     "C": f"=SUMIFS({C_INC},{C_DATE},$A{rr})",
     "D": f"=SUMIFS({C_XM},{C_DATE},$A{rr},{C_INC},1)",
     "E": f"=SUMIFS({C_GSR},{C_DATE},$A{rr},{C_INC},1)",
     "F": f"=D{rr}+E{rr}",
     "G": (f"=F{rr}-F{rr-1}" if k else "-"),
     "H": (f'=IF(F{rr-1}=0,"-",(F{rr}-F{rr-1})/F{rr-1})' if k else "-"),
     "I": f"=AVERAGE(F{max(DA1, rr-6)}:F{rr})",
     "J": (f'=IF(F{rr-7}=0,"-",(F{rr}-F{rr-7})/F{rr-7})' if k >= 7 else "-"),
     "K": f"=SUMIFS({C_DF},{C_DATE},$A{rr})",
     "L": f"=IFERROR(F{rr}/K{rr},0)",
     "M": f'=COUNTIFS({C_DATE},$A{rr},{C_XMC},"Yes",{C_INC},1)',
     "N": f'=M{rr}+COUNTIFS({C_DATE},$A{rr},{C_XMC},"No",{C_INC},1)',
     "O": f'=IF(N{rr}=0,"-",M{rr}/N{rr})',
     "P": f'=COUNTIFS({C_DATE},$A{rr},{C_ESIM},"Yes",{C_INC},1)',
     "Q": f'=P{rr}+COUNTIFS({C_DATE},$A{rr},{C_ESIM},"No",{C_INC},1)',
     "R": f'=IF(Q{rr}=0,"-",P{rr}/Q{rr})',
     "S": f"=SUMIFS({rg('P')},{C_DATE},$A{rr},{C_INC},1)",
     "T": f"=SUMIFS({rg('Q')},{C_DATE},$A{rr},{C_INC},1)",
     "U": f"=SUMIFS({rg('R')},{C_DATE},$A{rr},{C_INC},1)",
     "V": f"=SUMIFS({rg('S')},{C_DATE},$A{rr},{C_INC},1)",
     "W": f"=SUMIFS({rg('T')},{C_DATE},$A{rr},{C_INC},1)",
     "X": f'=COUNTIFS({C_DATE},$A{rr},{C_DEV},"<>No Device",{C_DEV},"<>",{C_INC},1)',
     "Y": f"=COUNTIFS({C_DATE},$A{rr},{C_INR},0,{C_INC},1)",
     "Z": f"=COUNTIFS({C_DATE},$A{rr},{C_DUP},1)"}
    fmts = {"H":"+0.0%;-0.0%;0.0%","I":"0.0","J":"+0.0%;-0.0%;0.0%","L":"0.00","O":"0.0%","R":"0.0%","G":"+#,##0;-#,##0;0"}
    for col, v in f.items(): put(da, f"{col}{rr}", v, F(10), fmt=fmts.get(col, "#,##0"))
tr = DAL + 1
put(da, f"A{tr}", "TOTAL", F(10, True, "FFFFFF"), NAVY); put(da, f"B{tr}", "", F(10), NAVY)
for col in "CDEFKMNPQSTUVWXYZ":
    put(da, f"{col}{tr}", f"=SUM({col}{DA1}:{col}{DAL})", F(10, True), LIGHT, "#,##0")
put(da, f"G{tr}", "", F(10), LIGHT); put(da, f"H{tr}", "", F(10), LIGHT); put(da, f"J{tr}", "", F(10), LIGHT)
put(da, f"I{tr}", f"=AVERAGE(F{DA1}:F{DAL})", F(10, True), LIGHT, "0.0")
put(da, f"L{tr}", f"=IFERROR(F{tr}/K{tr},0)", F(10, True), LIGHT, "0.00")
put(da, f"O{tr}", f'=IF(N{tr}=0,"-",M{tr}/N{tr})', F(10, True), LIGHT, "0.0%")
put(da, f"R{tr}", f'=IF(Q{tr}=0,"-",P{tr}/Q{tr})', F(10, True), LIGHT, "0.0%")
da.freeze_panes = "C4"
widths(da, {"A":11,"B":8,"C":11,"D":9,"E":9,"F":10,"G":10,"H":9,"I":9,"J":9,"K":9,"L":9,"M":8,"N":8,"O":9,"P":9,"Q":9,"R":9,"S":9,"T":8,"U":9,"V":9,"W":9,"X":9,"Y":11,"Z":11})
da.conditional_formatting.add(f"G{DA1}:G{DAL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
da.conditional_formatting.add(f"G{DA1}:G{DAL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
for col in "HJ":
    da.conditional_formatting.add(f"{col}{DA1}:{col}{DAL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN)))
    da.conditional_formatting.add(f"{col}{DA1}:{col}{DAL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED)))
da.conditional_formatting.add(f"F{DA1}:F{DAL}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
da.sheet_properties.tabColor = TEAL

# ================================================================= Daily_Leader
dl = wb.create_sheet("Daily_Leader")
NT = len(tls); TC1 = 2; TCL = TC1 + NT - 1; TOTC = TCL + 1; NRC = TOTC + 1
title(dl, "Leader Level – Day-over-Day", "Pick a metric in B3. Summary rows use the Dashboard date range; the daily matrix shows every day. Right-hand block = day-over-day change.", NRC)
put(dl, "A3", "Metric ▶", F(11, True), LIGHT); put(dl, "B3", METRICS[0], F(11, True, "0000FF"), INPUT)
dv_m = DataValidation(type="list", formula1="=Lists!$C$2:$C$5", allow_blank=False); dl.add_data_validation(dv_m); dv_m.add("B3")
put(dl, "A5", "Team Leader", F(10, True, "FFFFFF"), NAVY)
for j, t in enumerate(tls): put(dl, f"{L(TC1+j)}5", t, F(10, True, "FFFFFF"), NAVY).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
put(dl, f"{L(TOTC)}5", "Account Total", F(10, True, "FFFFFF"), NAVY); put(dl, f"{L(NRC)}5", "Not in Roster", F(10, True, "FFFFFF"), NAVY)
dl.row_dimensions[5].height = 30
srows = ["Headcount","Period Total","Avg / Day","Last Day (To Date)","Prior Day","DoD Δ","Rank (Period)"]
for i, lab in enumerate(srows): put(dl, f"A{6+i}", lab, F(10, True), LIGHT)
for j in range(NT + 2):
    col = L(TC1 + j); isTot = (j == NT); isNR = (j == NT + 1)
    tlcrit = [] if isTot else ([(C_TL, f"{col}$5")] if not isNR else [(C_TL, '"Not in Roster"')])
    hc = f"=COUNTIF({R_TL},{col}$5)" if not (isTot or isNR) else (f"=COUNTA({R_TL})" if isTot else "-")
    put(dl, f"{col}6", hc, F(10), fmt="#,##0")
    put(dl, f"{col}7", "=" + metric_choose("$B$3", [(C_DATE, f'">="&{D_FROM}'), (C_DATE, f'"<="&{D_TO}')] + tlcrit), F(10, True), fmt="#,##0")
    put(dl, f"{col}8", f"=IFERROR({col}7/{D_DAYS},0)", F(10), fmt="0.0")
    put(dl, f"{col}9", "=" + metric_choose("$B$3", [(C_DATE, D_TO)] + tlcrit), F(10), fmt="#,##0")
    put(dl, f"{col}10", "=" + metric_choose("$B$3", [(C_DATE, f"{D_TO}-1")] + tlcrit), F(10), fmt="#,##0")
    put(dl, f"{col}11", f"={col}9-{col}10", F(10, True), fmt="+#,##0;-#,##0;0")
    put(dl, f"{col}12", (f"=RANK({col}7,${L(TC1)}$7:${L(TCL)}$7)" if not (isTot or isNR) else "-"), F(10, True), fmt="0")
MR = 14; ML1 = MR + 1; MLL = ML1 + ND - 1
put(dl, f"A{MR}", "Date", F(10, True, "FFFFFF"), NAVY)
for j in range(NT + 2): put(dl, f"{L(TC1+j)}{MR}", f"={L(TC1+j)}$5", F(10, True, "FFFFFF"), NAVY).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
DC1 = NRC + 2
put(dl, f"{L(DC1-1)}{MR}", "", border=False)
put(dl, f"{L(DC1)}{MR}", "Date", F(10, True, "FFFFFF"), TEAL)
for j in range(NT + 1): put(dl, f"{L(DC1+1+j)}{MR}", f'="DoD Δ "&{L(TC1+j)}$5', F(10, True, "FFFFFF"), TEAL).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
dl.row_dimensions[MR].height = 30
for k, d in enumerate(dates):
    rr = ML1 + k
    put(dl, f"A{rr}", d, F(10, True), fmt="ddd dd-mmm")
    for j in range(NT + 2):
        col = L(TC1 + j)
        if j < NT: fml = "=" + metric_choose("$B$3", [(C_DATE, f"$A{rr}"), (C_TL, f"{col}${MR}")])
        elif j == NT: fml = "=" + metric_choose("$B$3", [(C_DATE, f"$A{rr}")])
        else: fml = f"={L(TOTC)}{rr}-SUM({L(TC1)}{rr}:{L(TCL)}{rr})"
        put(dl, f"{col}{rr}", fml, F(10, j >= NT), fmt="#,##0")
    put(dl, f"{L(DC1)}{rr}", f"=A{rr}", F(10, True), fmt="ddd dd-mmm")
    for j in range(NT + 1):
        col = L(TC1 + j)
        put(dl, f"{L(DC1+1+j)}{rr}", (f"={col}{rr}-{col}{rr-1}" if k else "-"), F(10), fmt="+#,##0;-#,##0;0")
dl.freeze_panes = f"B{ML1}"
dl.column_dimensions["A"].width = 16
for j in range(NT + 2): dl.column_dimensions[L(TC1 + j)].width = 13
dl.column_dimensions[L(DC1 - 1)].width = 2; dl.column_dimensions[L(DC1)].width = 13
for j in range(NT + 1): dl.column_dimensions[L(DC1 + 1 + j)].width = 13
dl.conditional_formatting.add(f"{L(TC1)}{ML1}:{L(TCL)}{MLL}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="63BE7B"))
dl.conditional_formatting.add(f"{L(DC1+1)}{ML1}:{L(DC1+NT+1)}{MLL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
dl.conditional_formatting.add(f"{L(DC1+1)}{ML1}:{L(DC1+NT+1)}{MLL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
dl.conditional_formatting.add(f"{L(TC1)}11:{L(NRC)}11", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
dl.conditional_formatting.add(f"{L(TC1)}11:{L(NRC)}11", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
dl.conditional_formatting.add(f"{L(TC1)}12:{L(TCL)}12", ColorScaleRule(start_type="min", start_color="63BE7B", end_type="max", end_color="F8696B"))
dl.sheet_properties.tabColor = TEAL

# ================================================================= Daily_Agent (team view)
dg = wb.create_sheet("Daily_Agent")
NA = 25; AC1 = 2; ACL = AC1 + NA - 1; ATOT = ACL + 1
title(dg, "Agent Level – Team Day-over-Day", "Pick a Team Leader in B3 and a metric in B4. Agents of that team load automatically. Summary rows use the Dashboard date range. Right-hand block = day-over-day change.", 14)
put(dg, "A3", "Team Leader ▶", F(11, True), LIGHT); put(dg, "B3", "Omar Mansour" if "Omar Mansour" in tls else tls[0], F(11, True, "0000FF"), INPUT)
put(dg, "A4", "Metric ▶", F(11, True), LIGHT); put(dg, "B4", METRICS[0], F(11, True, "0000FF"), INPUT)
dv_tl = DataValidation(type="list", formula1=f"=Lists!$A$3:$A${2+NT}", allow_blank=False); dg.add_data_validation(dv_tl); dv_tl.add("B3")
dv_m2 = DataValidation(type="list", formula1="=Lists!$C$2:$C$5", allow_blank=False); dg.add_data_validation(dv_m2); dv_m2.add("B4")
put(dg, "A6", "Agent", F(10, True, "FFFFFF"), NAVY); put(dg, "A7", "HRID", F(10, True), LIGHT); put(dg, "A8", "Phase", F(10, True), LIGHT)
arows = ["Period Total","Avg / Day","Last Day (To Date)","Prior Day","DoD Δ","Rank in Team","Active Days (Period)"]
for i, lab in enumerate(arows): put(dg, f"A{9+i}", lab, F(10, True), LIGHT)
put(dg, f"{L(ATOT)}6", "Team Total", F(10, True, "FFFFFF"), NAVY)
for j in range(NA):
    col = L(AC1 + j); n = j + 1
    put(dg, f"{col}6", f'=IFERROR(INDEX({R_NAME},MATCH($B$3&"|"&{n},{R_KEY},0)),"")', F(10, True, "FFFFFF"), NAVY).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    put(dg, f"{col}7", f'=IFERROR(INDEX({R_HR},MATCH($B$3&"|"&{n},{R_KEY},0)),"")', F(10))
    put(dg, f"{col}8", f'=IFERROR(INDEX({R_PHASE},MATCH($B$3&"|"&{n},{R_KEY},0)),"")', F(10))
    g = f'IF({col}$7="","",'
    put(dg, f"{col}9", f"={g}" + metric_choose("$B$4", [(C_DATE, f'">="&{D_FROM}'), (C_DATE, f'"<="&{D_TO}'), (C_HR, f"{col}$7")]) + ")", F(10, True), fmt="#,##0")
    put(dg, f"{col}10", f'={g}{col}9/{D_DAYS})', F(10), fmt="0.0")
    put(dg, f"{col}11", f"={g}" + metric_choose("$B$4", [(C_DATE, D_TO), (C_HR, f"{col}$7")]) + ")", F(10), fmt="#,##0")
    put(dg, f"{col}12", f"={g}" + metric_choose("$B$4", [(C_DATE, f"{D_TO}-1"), (C_HR, f"{col}$7")]) + ")", F(10), fmt="#,##0")
    put(dg, f"{col}13", f"={g}{col}11-{col}12)", F(10, True), fmt="+#,##0;-#,##0;0")
    put(dg, f"{col}14", f"={g}RANK({col}9,${L(AC1)}$9:${L(ACL)}$9))", F(10, True), fmt="0")
    put(dg, f"{col}15", f'={g}SUMIFS({C_DF},{C_HR},{col}$7,{C_DATE},">="&{D_FROM},{C_DATE},"<="&{D_TO}))', F(10), fmt="0")
for rr_ in range(7, 16): put(dg, f"{L(ATOT)}{rr_}", (f"=SUM({L(AC1)}{rr_}:{L(ACL)}{rr_})" if rr_ in (9, 11, 12, 13) else ""), F(10, True), LIGHT, "#,##0")
dg.row_dimensions[6].height = 45
GR = 17; GL1 = GR + 1; GLL = GL1 + ND - 1
put(dg, f"A{GR}", "Date", F(10, True, "FFFFFF"), NAVY)
for j in range(NA): put(dg, f"{L(AC1+j)}{GR}", f"={L(AC1+j)}$6", F(10, True, "FFFFFF"), NAVY).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
put(dg, f"{L(ATOT)}{GR}", "Team Total", F(10, True, "FFFFFF"), NAVY)
GD1 = ATOT + 2
put(dg, f"{L(GD1)}{GR}", "Date", F(10, True, "FFFFFF"), TEAL)
for j in range(NA + 1): put(dg, f"{L(GD1+1+j)}{GR}", f'=IF({L(AC1+j)}$6="","","DoD Δ "&{L(AC1+j)}$6)', F(10, True, "FFFFFF"), TEAL).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
dg.row_dimensions[GR].height = 45
for k, d in enumerate(dates):
    rr = GL1 + k
    put(dg, f"A{rr}", d, F(10, True), fmt="ddd dd-mmm")
    for j in range(NA):
        col = L(AC1 + j)
        put(dg, f"{col}{rr}", f'=IF({col}$7="","",' + metric_choose("$B$4", [(C_DATE, f"$A{rr}"), (C_HR, f"{col}$7")]) + ")", F(10), fmt="#,##0")
    put(dg, f"{L(ATOT)}{rr}", f"=SUM({L(AC1)}{rr}:{L(ACL)}{rr})", F(10, True), fmt="#,##0")
    put(dg, f"{L(GD1)}{rr}", f"=A{rr}", F(10, True), fmt="ddd dd-mmm")
    for j in range(NA + 1):
        col = L(AC1 + j)
        put(dg, f"{L(GD1+1+j)}{rr}", (f'=IF({col}{rr}="","",{col}{rr}-{col}{rr-1})' if k else "-"), F(10), fmt="+#,##0;-#,##0;0")
dg.freeze_panes = f"B{GL1}"
dg.column_dimensions["A"].width = 20
for j in range(NA + 1): dg.column_dimensions[L(AC1 + j)].width = 13
dg.column_dimensions[L(GD1 - 1)].width = 2; dg.column_dimensions[L(GD1)].width = 13
for j in range(NA + 1): dg.column_dimensions[L(GD1 + 1 + j)].width = 13
dg.conditional_formatting.add(f"{L(AC1)}{GL1}:{L(ACL)}{GLL}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color="63BE7B"))
dg.conditional_formatting.add(f"{L(GD1+1)}{GL1}:{L(GD1+NA+1)}{GLL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
dg.conditional_formatting.add(f"{L(GD1+1)}{GL1}:{L(GD1+NA+1)}{GLL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
dg.conditional_formatting.add(f"{L(AC1)}13:{L(ATOT)}13", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
dg.conditional_formatting.add(f"{L(AC1)}13:{L(ATOT)}13", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
dg.sheet_properties.tabColor = TEAL

# ================================================================= Agent_Summary
ag = wb.create_sheet("Agent_Summary")
title(ag, "Agent Summary – all roster agents (selected period & scope from Dashboard)", None, 22)
put(ag, "A2", "Data through ▶", F(10, True), LIGHT); put(ag, "B2", f"=MAX({C_DATE})", F(10, True), fmt="dd-mmm-yy")
put(ag, "C2", "Period ▶", F(10, True), LIGHT); put(ag, "D2", f'=TEXT({D_FROM},"dd-mmm")&" to "&TEXT({D_TO},"dd-mmm")', F(10, True))
put(ag, "E2", "Use the filter arrows to sort / filter. In Scope = 1 when the agent matches the Dashboard Ops Manager / Team Leader selection.", F(9, False, "595959", True), border=False)
acols = ["HRID","Agent Name","Team Leader","Ops Manager","Phase","Tenure Bucket","Wave","Submissions","XM Lines","GSR Units","Total Units","Active Days","Units / Active Day","Prior Period Units","Δ vs Prior","Last Sale Date","Days Since Last Sale","Rank (Scope)","Rank (Team)","In Scope","ScoreHigh","ScoreLow","LastSaleSerial"]
header_row(ag, 4, 1, acols); ag.row_dimensions[4].height = 32
for i, x in r.iterrows():
    rr = AS1 + i
    ri = i + 2
    put(ag, f"A{rr}", x["HRID"], F(10)); put(ag, f"B{rr}", x["Agent Name"], F(10, True)); put(ag, f"C{rr}", x["Team Leader"], F(10)); put(ag, f"D{rr}", x["Ops Manager"], F(10))
    put(ag, f"E{rr}", x["Phase"], F(10)); put(ag, f"F{rr}", x["Tenure Bucket"], F(10)); put(ag, f"G{rr}", x["Wave"], F(10))
    per = f'{C_DATE},">="&{D_FROM},{C_DATE},"<="&{D_TO}'
    f = {"H": f"=SUMIFS({C_INC},{C_HR},$A{rr},{per})",
         "I": f"=SUMIFS({C_XM},{C_HR},$A{rr},{per},{C_INC},1)",
         "J": f"=SUMIFS({C_GSR},{C_HR},$A{rr},{per},{C_INC},1)",
         "K": f"=SUMIFS({C_U},{C_HR},$A{rr},{per},{C_INC},1)",
         "L": f"=SUMIFS({C_DF},{C_HR},$A{rr},{per})",
         "M": f"=IFERROR(K{rr}/L{rr},0)",
         "N": f'=SUMIFS({C_U},{C_HR},$A{rr},{C_DATE},">="&{D_PFROM},{C_DATE},"<="&{D_PTO},{C_INC},1)',
         "O": f"=K{rr}-N{rr}",
         "P": f'=IF(W{rr}=0,"-",W{rr})',
         "Q": f'=IF(W{rr}=0,"-",$B$2-W{rr})',
         "R": f'=IF(T{rr}=1,COUNTIFS($T${AS1}:$T${ASL},1,$K${AS1}:$K${ASL},">"&K{rr})+1,"-")',
         "S": f'=COUNTIFS($C${AS1}:$C${ASL},C{rr},$K${AS1}:$K${ASL},">"&K{rr})+1',
         "T": f'=IF(AND(OR({D_OM}="All",D{rr}={D_OM}),OR({D_TL}="All",C{rr}={D_TL})),1,0)',
         "U": f"=IF(T{rr}=1,K{rr}+ROW()/1000000,-1)",
         "V": f'=IF(AND(T{rr}=1,E{rr}="Production"),K{rr}+ROW()/1000000,1000000000)',
         "W": f"=SUMPRODUCT(MAX(({C_HR}=A{rr})*{C_DATE}))"}
    fm = {"M":"0.00","O":"+#,##0;-#,##0;0","P":"dd-mmm-yy","Q":"0","R":"0","S":"0","T":"0","U":"0.000000","V":"0.000000","W":"0"}
    for col, v in f.items(): put(ag, f"{col}{rr}", v, F(10, col == "K"), fmt=fm.get(col, "#,##0"))
ag.freeze_panes = "C5"; ag.auto_filter.ref = f"A4:W{ASL}"
widths(ag, {"A":12,"B":32,"C":17,"D":16,"E":11,"F":10,"G":12,"H":11,"I":9,"J":9,"K":10,"L":9,"M":10,"N":11,"O":9,"P":11,"Q":10,"R":9,"S":9,"T":8,"U":10,"V":12,"W":10})
ag.conditional_formatting.add(f"K{AS1}:K{ASL}", DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
ag.conditional_formatting.add(f"O{AS1}:O{ASL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
ag.conditional_formatting.add(f"O{AS1}:O{ASL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
ag.conditional_formatting.add(f"Q{AS1}:Q{ASL}", CellIsRule(operator="greaterThan", formula=["7"], fill=fill("F8CBAD")))
for col in "UVW": ag.column_dimensions[col].hidden = True
ag.sheet_properties.tabColor = TEAL

# ================================================================= Agent_Trend
at = wb.create_sheet("Agent_Trend")
title(at, "Single Agent – Daily Trend vs Team", "Pick an agent in B3 (type to search in the dropdown). Team Avg / Agent = team units that day ÷ team headcount.", 20)
put(at, "A3", "Agent ▶", F(11, True), LIGHT); at.merge_cells("B3:D3")
top_hr = s[s.in_roster == 1].groupby("hr")["units"].sum().idxmax()
put(at, "B3", rmap.loc[top_hr, "Label"], F(11, True, "0000FF"), INPUT)
dv_ag = DataValidation(type="list", formula1=f"=Lists!$D$2:$D${1+RN}", allow_blank=False); at.add_data_validation(dv_ag); dv_ag.add("B3")
info = [("HRID", f'=IFERROR(INDEX({R_HR},MATCH($B$3,{R_LABEL},0)),"")'), ("Team Leader", f'=IFERROR(INDEX({R_TL},MATCH($B$3,{R_LABEL},0)),"")'),
        ("Ops Manager", f'=IFERROR(INDEX({R_OM},MATCH($B$3,{R_LABEL},0)),"")'), ("Phase", f'=IFERROR(INDEX({R_PHASE},MATCH($B$3,{R_LABEL},0)),"")'),
        ("Period Units (Dashboard range)", f'=IFERROR(INDEX({asr("K")},MATCH($B$4,{asr("A")},0)),0)'), ("Rank in Team (Period)", f'=IFERROR(INDEX({asr("S")},MATCH($B$4,{asr("A")},0)),"-")'),
        ("All-time Units", f"=SUMIFS({C_U},{C_HR},$B$4,{C_INC},1)"), ("Last Sale Date", f'=IFERROR(INDEX({asr("P")},MATCH($B$4,{asr("A")},0)),"-")')]
for i, (lab, fml) in enumerate(info):
    rr = 4 + i; put(at, f"A{rr}", lab, F(10, True), LIGHT); at.merge_cells(f"B{rr}:D{rr}"); put(at, f"B{rr}", fml, F(10, True), fmt=("dd-mmm-yy" if lab == "Last Sale Date" else "#,##0"))
TR = 13; TL1 = TR + 1; TLL = TL1 + ND - 1
header_row(at, TR, 1, ["Date","Submissions","XM Lines","GSR Units","Total Units","DoD Δ","7-Day Avg","Cumulative Units","Team Avg / Agent","vs Team Avg"]); at.row_dimensions[TR].height = 30
for k, d in enumerate(dates):
    rr = TL1 + k
    put(at, f"A{rr}", d, F(10, True), fmt="ddd dd-mmm")
    f = {"B": f"=SUMIFS({C_INC},{C_HR},$B$4,{C_DATE},$A{rr})", "C": f"=SUMIFS({C_XM},{C_HR},$B$4,{C_DATE},$A{rr},{C_INC},1)",
         "D": f"=SUMIFS({C_GSR},{C_HR},$B$4,{C_DATE},$A{rr},{C_INC},1)", "E": f"=C{rr}+D{rr}", "F": (f"=E{rr}-E{rr-1}" if k else "-"),
         "G": f"=AVERAGE(E{max(TL1, rr-6)}:E{rr})", "H": f"=SUM($E${TL1}:E{rr})",
         "I": f"=IFERROR(SUMIFS({C_U},{C_DATE},$A{rr},{C_TL},$B$5,{C_INC},1)/COUNTIF({R_TL},$B$5),0)", "J": f"=E{rr}-I{rr}"}
    fm = {"F":"+#,##0;-#,##0;0","G":"0.0","I":"0.00","J":"+0.00;-0.00;0.00"}
    for col, v in f.items(): put(at, f"{col}{rr}", v, F(10, col == "E"), fmt=fm.get(col, "#,##0"))
at.freeze_panes = f"B{TL1}"
widths(at, {"A":26,"B":12,"C":10,"D":10,"E":11,"F":9,"G":10,"H":12,"I":13,"J":12})
for col in "FJ":
    at.conditional_formatting.add(f"{col}{TL1}:{col}{TLL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
    at.conditional_formatting.add(f"{col}{TL1}:{col}{TLL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
ch = LineChart(); ch.title = "Agent daily units vs team average per agent"; ch.height = 9; ch.width = 22; ch.style = 10
ch.add_data(Reference(at, min_col=5, min_row=TR, max_row=TLL), titles_from_data=True)
ch.add_data(Reference(at, min_col=9, min_row=TR, max_row=TLL), titles_from_data=True)
ch.set_categories(Reference(at, min_col=1, min_row=TL1, max_row=TLL)); ch.y_axis.title = "Units"; ch.x_axis.number_format = "dd-mmm"; ch.legend.position = "b"
at.add_chart(ch, "L4")
at.sheet_properties.tabColor = TEAL

# ================================================================= Dashboard
db = wb.create_sheet("Dashboard", 0)
db.sheet_view.showGridLines = False
title(db, "Cairo Repair TSX – Sales Performance Dashboard", f"Source: Sales Tracker (Microsoft Forms) {min(dates):%d %b %Y} – {MAXD:%d %b %Y} + TSX Roster. Change the yellow cells; every table, rank and chart recalculates.", 14)
section(db, 4, "CONTROL PANEL", 1, 3)
ctrl = [("From Date", DEF_FROM, "dd-mmm-yy"), ("To Date", DEF_TO, "dd-mmm-yy"), ("Ops Manager", "All", None), ("Team Leader", "All", None), ("Exclude Duplicates?", "Yes", None)]
for i, (lab, v, fmt) in enumerate(ctrl):
    rr = 5 + i; db.merge_cells(f"A{rr}:B{rr}"); put(db, f"A{rr}", lab, F(11, True), LIGHT); put(db, f"C{rr}", v, F(11, True, "0000FF"), INPUT, fmt)
db.merge_cells("A10:B10"); put(db, "A10", "Days in Period", F(10, True), GREY); put(db, "C10", "=C6-C5+1", F(10, True), GREY, "0")
db.merge_cells("A11:B11"); put(db, "A11", "Prior Period (same length)", F(10, True), GREY); put(db, "C11", '=TEXT(C12,"dd-mmm")&" → "&TEXT(C13,"dd-mmm")', F(10, True), GREY)
helpers = [("Prior From", "=C5-C10", "dd-mmm-yy"), ("Prior To", "=C5-1", "dd-mmm-yy"), ("TL criterion", '=IF(C8="All","*",C8)', None), ("OM criterion", '=IF(C7="All","*",C7)', None)]
for i, (lab, v, fmt) in enumerate(helpers):
    rr = 12 + i; db.merge_cells(f"A{rr}:B{rr}"); put(db, f"A{rr}", lab, F(9, False, "7F7F7F"), GREY); put(db, f"C{rr}", v, F(9, False, "7F7F7F"), GREY, fmt)
dv1 = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${2+len(oms)}", allow_blank=False); db.add_data_validation(dv1); dv1.add("C7")
dv2 = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${2+NT}", allow_blank=False); db.add_data_validation(dv2); dv2.add("C8")
dv3 = DataValidation(type="list", formula1="=Lists!$E$2:$E$3", allow_blank=False); db.add_data_validation(dv3); dv3.add("C9")
dv4 = DataValidation(type="date", operator="between", formula1=f"DATE({min(dates).year},{min(dates).month},{min(dates).day})", formula2=f"DATE({MAXD.year},{MAXD.month},{MAXD.day})", allow_blank=False)
db.add_data_validation(dv4); dv4.add("C5"); dv4.add("C6")

def sc(sumcol, frm, to, extra=""):
    return f'SUMIFS({sumcol},{C_DATE},">="&{frm},{C_DATE},"<="&{to},{C_TL},{D_TLC},{C_OM},{D_OMC},{C_INC},1{extra})'
def cc(frm, to, extra):
    return f'COUNTIFS({C_DATE},">="&{frm},{C_DATE},"<="&{to},{C_TL},{D_TLC},{C_OM},{D_OMC},{C_INC},1{extra})'
cur = (D_FROM, D_TO); pri = (D_PFROM, D_PTO)
ACT_CUR = f'COUNTIFS({asr("T")},1,{asr("K")},">0")'; ACT_PRI = f'COUNTIFS({asr("T")},1,{asr("N")},">0")'; HC_SC = f'COUNTIF({asr("T")},1)'
tiles = [
 ("Total Units", f"={sc(C_U,*cur)}", f"={sc(C_U,*pri)}", "#,##0", "+#,##0;-#,##0;0"),
 ("XM Lines", f"={sc(C_XM,*cur)}", f"={sc(C_XM,*pri)}", "#,##0", "+#,##0;-#,##0;0"),
 ("GSR Units", f"={sc(C_GSR,*cur)}", f"={sc(C_GSR,*pri)}", "#,##0", "+#,##0;-#,##0;0"),
 ("Submissions", f"={sc(C_INC,*cur)}", f"={sc(C_INC,*pri)}", "#,##0", "+#,##0;-#,##0;0"),
 ("Active Sellers", f"={ACT_CUR}", f"={ACT_PRI}", "#,##0", "+#,##0;-#,##0;0"),
 ("Units / Active Seller", f"=IFERROR({sc(C_U,*cur)}/{ACT_CUR},0)", f"=IFERROR({sc(C_U,*pri)}/{ACT_PRI},0)", "0.00", "+0.00;-0.00;0.00"),
 ("Participation % (of HC in scope)", f"=IFERROR({ACT_CUR}/{HC_SC},0)", f"=IFERROR({ACT_PRI}/{HC_SC},0)", "0.0%", "+0.0%;-0.0%;0.0%"),
 ("XMC Attach %", f'=IFERROR({cc(*cur, "," + C_XMC + ",\"Yes\"")}/({cc(*cur, "," + C_XMC + ",\"Yes\"")}+{cc(*cur, "," + C_XMC + ",\"No\"")}),0)',
                  f'=IFERROR({cc(*pri, "," + C_XMC + ",\"Yes\"")}/({cc(*pri, "," + C_XMC + ",\"Yes\"")}+{cc(*pri, "," + C_XMC + ",\"No\"")}),0)', "0.0%", "+0.0%;-0.0%;0.0%"),
 ("E-SIM Activation %", f'=IFERROR({cc(*cur, "," + C_ESIM + ",\"Yes\"")}/({cc(*cur, "," + C_ESIM + ",\"Yes\"")}+{cc(*cur, "," + C_ESIM + ",\"No\"")}),0)',
                  f'=IFERROR({cc(*pri, "," + C_ESIM + ",\"Yes\"")}/({cc(*pri, "," + C_ESIM + ",\"Yes\"")}+{cc(*pri, "," + C_ESIM + ",\"No\"")}),0)', "0.0%", "+0.0%;-0.0%;0.0%"),
 ("Avg Units / Day", f"=IFERROR({sc(C_U,*cur)}/{D_DAYS},0)", f"=IFERROR({sc(C_U,*pri)}/{D_DAYS},0)", "0.0", "+0.0;-0.0;0.0"),
]
for i, (lab, fcur, fpri, vf, df) in enumerate(tiles):
    r0 = 4 if i < 5 else 9; c0 = 5 + 2 * (i % 5); cl = L(c0); cr = L(c0 + 1)
    put(db, f"{cl}{r0}", lab, F(10, True, "FFFFFF"), NAVY); db[f"{cr}{r0}"].fill = fill(NAVY); db[f"{cr}{r0}"].border = BORDER
    put(db, f"{cl}{r0+1}", fcur, F(20, True, NAVY), GREY, vf); db[f"{cr}{r0+1}"].fill = fill(GREY); db[f"{cr}{r0+1}"].border = BORDER
    put(db, f"{cl}{r0+2}", f"={cl}{r0+1}-({fpri[1:]})", F(11, True), GREY, df); db[f"{cr}{r0+2}"].fill = fill(GREY); db[f"{cr}{r0+2}"].border = BORDER
    put(db, f"{cl}{r0+3}", "vs prior period", F(8, False, "7F7F7F", True), GREY); db[f"{cr}{r0+3}"].fill = fill(GREY); db[f"{cr}{r0+3}"].border = BORDER
    for rr in range(r0, r0 + 4): db.merge_cells(f"{cl}{rr}:{cr}{rr}")
    db.row_dimensions[r0 + 1].height = 30
    db.conditional_formatting.add(f"{cl}{r0+2}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
    db.conditional_formatting.add(f"{cl}{r0+2}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))

# --- account DoD mini table (last 8 days ending To Date)
section(db, 17, "ACCOUNT LEVEL – LAST 8 DAYS DAY-OVER-DAY (whole account, ending at To Date)")
header_row(db, 18, 1, ["Date","Weekday","Submissions","XM Lines","GSR Units","Total Units","DoD Δ","DoD %","7-Day Avg","WoW %","Active Sellers","Units / Seller","XMC Attach %","E-SIM Act. %"], LIGHT, NAVY, 10); db.row_dimensions[18].height = 30
dacol = {"B":"B","C":"C","D":"D","E":"E","F":"F","G":"G","H":"H","I":"I","J":"J","K":"K","L":"L","M":"O","N":"R"}
for k in range(8):
    rr = 19 + k
    put(db, f"A{rr}", f"=$C$6-{7-k}", F(10, True), fmt="dd-mmm-yy")
    for dcol, scol in dacol.items():
        put(db, f"{dcol}{rr}", f'=IFERROR(INDEX(Daily_Account!${scol}${DA1}:${scol}${DAL},MATCH($A{rr},Daily_Account!$A${DA1}:$A${DAL},0)),"-")', F(10, dcol == "F"), fmt={"H":"+0.0%;-0.0%;0.0%","J":"+0.0%;-0.0%;0.0%","I":"0.0","L":"0.00","M":"0.0%","N":"0.0%","G":"+#,##0;-#,##0;0"}.get(dcol, "#,##0"))
for col in "GHJ":
    db.conditional_formatting.add(f"{col}19:{col}26", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
    db.conditional_formatting.add(f"{col}19:{col}26", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))

# --- leader scorecard
section(db, 28, "LEADER SCORECARD – selected period (Ops Manager filter applies; Team Leader filter ignored here so leaders stay comparable)")
header_row(db, 29, 1, ["Rank","Team Leader","Ops Manager","HC","Submissions","XM Lines","GSR Units","Total Units","Units / HC","Active Sellers","Participation %","Prior Period Units","Δ vs Prior","Δ %"], LIGHT, NAVY, 10); db.row_dimensions[29].height = 30
LS1 = 30; LSL = LS1 + NT - 1
for j, t in enumerate(tls):
    rr = LS1 + j; om = " / ".join(r[r["Team Leader"] == t]["Ops Manager"].value_counts().index.tolist())
    per = f'{C_DATE},">="&$C$5,{C_DATE},"<="&$C$6'; pper = f'{C_DATE},">="&$C$12,{C_DATE},"<="&$C$13'
    f = {"A": f'=IF(H{rr}=0,"-",RANK(H{rr},$H${LS1}:$H${LSL}))', "B": t, "C": om, "D": f"=COUNTIF({R_TL},$B{rr})",
         "E": f"=SUMIFS({C_INC},{per},{C_TL},$B{rr},{C_OM},$C$15)", "F": f"=SUMIFS({C_XM},{per},{C_TL},$B{rr},{C_OM},$C$15,{C_INC},1)",
         "G": f"=SUMIFS({C_GSR},{per},{C_TL},$B{rr},{C_OM},$C$15,{C_INC},1)", "H": f"=F{rr}+G{rr}", "I": f"=IFERROR(H{rr}/D{rr},0)",
         "J": f'=COUNTIFS({asr("C")},$B{rr},{asr("D")},$C$15,{asr("K")},">0")', "K": f"=IFERROR(J{rr}/D{rr},0)",
         "L": f"=SUMIFS({C_U},{pper},{C_TL},$B{rr},{C_OM},$C$15,{C_INC},1)", "M": f"=H{rr}-L{rr}", "N": f'=IF(L{rr}=0,"-",(H{rr}-L{rr})/L{rr})'}
    fm = {"I":"0.00","K":"0.0%","M":"+#,##0;-#,##0;0","N":"+0.0%;-0.0%;0.0%","A":"0"}
    for col, v in f.items(): put(db, f"{col}{rr}", v, F(10, col in "ABH"), fmt=fm.get(col, "#,##0"))
tr2 = LSL + 1
put(db, f"A{tr2}", "", fl=NAVY); put(db, f"B{tr2}", "ACCOUNT TOTAL (roster teams)", F(10, True, "FFFFFF"), NAVY); put(db, f"C{tr2}", "", fl=NAVY)
for col in "DEFGHJLM": put(db, f"{col}{tr2}", f"=SUM({col}{LS1}:{col}{LSL})", F(10, True), LIGHT, "+#,##0;-#,##0;0" if col == "M" else "#,##0")
put(db, f"I{tr2}", f"=IFERROR(H{tr2}/D{tr2},0)", F(10, True), LIGHT, "0.00"); put(db, f"K{tr2}", f"=IFERROR(J{tr2}/D{tr2},0)", F(10, True), LIGHT, "0.0%"); put(db, f"N{tr2}", f'=IF(L{tr2}=0,"-",(H{tr2}-L{tr2})/L{tr2})', F(10, True), LIGHT, "+0.0%;-0.0%;0.0%")
nr = tr2 + 1
put(db, f"A{nr}", "", fl=GREY); db.merge_cells(f"B{nr}:C{nr}"); put(db, f"B{nr}", "Not in Roster (unmatched HRIDs)", F(10, True, "7F7F7F"), GREY)
put(db, f"D{nr}", "-", F(10), GREY); put(db, f"E{nr}", f'=SUMIFS({C_INC},{C_DATE},">="&$C$5,{C_DATE},"<="&$C$6,{C_INR},0)', F(10), GREY, "#,##0")
put(db, f"F{nr}", f'=SUMIFS({C_XM},{C_DATE},">="&$C$5,{C_DATE},"<="&$C$6,{C_INR},0,{C_INC},1)', F(10), GREY, "#,##0")
put(db, f"G{nr}", f'=SUMIFS({C_GSR},{C_DATE},">="&$C$5,{C_DATE},"<="&$C$6,{C_INR},0,{C_INC},1)', F(10), GREY, "#,##0"); put(db, f"H{nr}", f"=F{nr}+G{nr}", F(10, True), GREY, "#,##0")
for col in "IJK": put(db, f"{col}{nr}", "-", F(10), GREY)
put(db, f"L{nr}", f'=SUMIFS({C_U},{C_DATE},">="&$C$12,{C_DATE},"<="&$C$13,{C_INR},0,{C_INC},1)', F(10), GREY, "#,##0"); put(db, f"M{nr}", f"=H{nr}-L{nr}", F(10), GREY, "+#,##0;-#,##0;0"); put(db, f"N{nr}", "-", F(10), GREY)
db.conditional_formatting.add(f"I{LS1}:I{LSL}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
db.conditional_formatting.add(f"K{LS1}:K{LSL}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
for col in "MN":
    db.conditional_formatting.add(f"{col}{LS1}:{col}{LSL}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
    db.conditional_formatting.add(f"{col}{LS1}:{col}{LSL}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))
db.conditional_formatting.add(f"H{LS1}:H{LSL}", DataBarRule(start_type="num", start_value=0, end_type="max", color="2E75B6"))

# --- ops manager block
OM0 = nr + 2
section(db, OM0, "OPS MANAGER VIEW – selected period")
header_row(db, OM0 + 1, 1, ["Rank","Ops Manager","Teams","HC","Submissions","XM Lines","GSR Units","Total Units","Units / HC","Active Sellers","Participation %","Prior Period Units","Δ vs Prior","Δ %"], LIGHT, NAVY, 10); db.row_dimensions[OM0 + 1].height = 30
OM1 = OM0 + 2; OML = OM1 + len(oms) - 1
for j, o in enumerate(oms):
    rr = OM1 + j; per = f'{C_DATE},">="&$C$5,{C_DATE},"<="&$C$6'; pper = f'{C_DATE},">="&$C$12,{C_DATE},"<="&$C$13'
    f = {"A": f'=IF(H{rr}=0,"-",RANK(H{rr},$H${OM1}:$H${OML}))', "B": o, "C": f'=SUMPRODUCT(({R_OM}=$B{rr})/COUNTIFS({R_TL},{R_TL},{R_OM},{R_OM}))', "D": f"=COUNTIF({R_OM},$B{rr})",
         "E": f"=SUMIFS({C_INC},{per},{C_OM},$B{rr})", "F": f"=SUMIFS({C_XM},{per},{C_OM},$B{rr},{C_INC},1)", "G": f"=SUMIFS({C_GSR},{per},{C_OM},$B{rr},{C_INC},1)",
         "H": f"=F{rr}+G{rr}", "I": f"=IFERROR(H{rr}/D{rr},0)", "J": f'=COUNTIFS({asr("D")},$B{rr},{asr("K")},">0")', "K": f"=IFERROR(J{rr}/D{rr},0)",
         "L": f"=SUMIFS({C_U},{pper},{C_OM},$B{rr},{C_INC},1)", "M": f"=H{rr}-L{rr}", "N": f'=IF(L{rr}=0,"-",(H{rr}-L{rr})/L{rr})'}
    fm = {"I":"0.00","K":"0.0%","M":"+#,##0;-#,##0;0","N":"+0.0%;-0.0%;0.0%","A":"0","C":"0"}
    for col, v in f.items(): put(db, f"{col}{rr}", v, F(10, col in "ABH"), fmt=fm.get(col, "#,##0"))
for col in "MN":
    db.conditional_formatting.add(f"{col}{OM1}:{col}{OML}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name="Calibri", color=GREEN, bold=True)))
    db.conditional_formatting.add(f"{col}{OM1}:{col}{OML}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Calibri", color=RED, bold=True)))

# --- top / bottom agents
TB0 = OML + 2
section(db, TB0, "TOP 10 AGENTS – selected period & scope", 1, 6); section(db, TB0, "BOTTOM 10 PRODUCTION AGENTS – selected period & scope (coaching list)", 8, 14)
header_row(db, TB0 + 1, 1, ["#","Agent","Team Leader","Units","XM","GSR"], LIGHT, NAVY, 10)
header_row(db, TB0 + 1, 8, ["#","Agent","Team Leader","Units","Active Days","Last Sale","Days Since"], LIGHT, NAVY, 10)
for k in range(10):
    rr = TB0 + 2 + k
    hi = f"LARGE({asr('U')},{k+1})"; lo = f"SMALL({asr('V')},{k+1})"
    mh = f"MATCH({hi},{asr('U')},0)"; ml = f"MATCH({lo},{asr('V')},0)"
    put(db, f"A{rr}", k + 1, F(10, True), fmt="0")
    put(db, f"B{rr}", f'=IF({hi}<0,"",INDEX({asr("B")},{mh}))', F(10, True)); put(db, f"C{rr}", f'=IF({hi}<0,"",INDEX({asr("C")},{mh}))', F(10))
    put(db, f"D{rr}", f'=IF({hi}<0,"",INDEX({asr("K")},{mh}))', F(10, True), fmt="#,##0"); put(db, f"E{rr}", f'=IF({hi}<0,"",INDEX({asr("I")},{mh}))', F(10), fmt="#,##0"); put(db, f"F{rr}", f'=IF({hi}<0,"",INDEX({asr("J")},{mh}))', F(10), fmt="#,##0")
    put(db, f"H{rr}", k + 1, F(10, True), fmt="0")
    put(db, f"I{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("B")},{ml}))', F(10, True)); put(db, f"J{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("C")},{ml}))', F(10))
    put(db, f"K{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("K")},{ml}))', F(10, True), fmt="#,##0"); put(db, f"L{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("L")},{ml}))', F(10), fmt="#,##0")
    put(db, f"M{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("P")},{ml}))', F(10), fmt="dd-mmm-yy"); put(db, f"N{rr}", f'=IF({lo}>=1000000000,"",INDEX({asr("Q")},{ml}))', F(10), fmt="0")
db.conditional_formatting.add(f"D{TB0+2}:D{TB0+11}", DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
db.conditional_formatting.add(f"K{TB0+2}:K{TB0+11}", CellIsRule(operator="equal", formula=["0"], fill=fill("F8CBAD")))

# --- charts
CH0 = TB0 + 13
section(db, CH0, "CHARTS – account trend (all days), leader comparison (selected period), daily mix")
c1 = LineChart(); c1.title = "Account – daily Total Units with 7-day average"; c1.height = 9; c1.width = 22; c1.style = 10
c1.add_data(Reference(da, min_col=6, min_row=3, max_row=DAL), titles_from_data=True); c1.add_data(Reference(da, min_col=9, min_row=3, max_row=DAL), titles_from_data=True)
c1.set_categories(Reference(da, min_col=1, min_row=DA1, max_row=DAL)); c1.y_axis.title = "Units"; c1.x_axis.number_format = "dd-mmm"; c1.legend.position = "b"
db.add_chart(c1, f"A{CH0+1}")
c2 = BarChart(); c2.type = "bar"; c2.title = "Leaders – Total Units (selected period)"; c2.height = 9; c2.width = 16; c2.style = 10
c2.add_data(Reference(db, min_col=8, min_row=29, max_row=LSL), titles_from_data=True); c2.set_categories(Reference(db, min_col=2, min_row=LS1, max_row=LSL)); c2.legend = None
c2.dataLabels = DataLabelList(); c2.dataLabels.showVal = True; c2.dataLabels.showSerName = False; c2.dataLabels.showCatName = False; c2.dataLabels.showLegendKey = False; c2.gapWidth = 40
db.add_chart(c2, f"I{CH0+1}")
c3 = BarChart(); c3.type = "col"; c3.grouping = "stacked"; c3.overlap = 100; c3.title = "Account – daily XM Lines vs GSR Units"; c3.height = 9; c3.width = 22; c3.style = 10
c3.gapWidth = 30; c3.legend.position = "b"; c3.add_data(Reference(da, min_col=4, min_row=3, max_col=5, max_row=DAL), titles_from_data=True); c3.set_categories(Reference(da, min_col=1, min_row=DA1, max_row=DAL)); c3.x_axis.number_format = "dd-mmm"
db.add_chart(c3, f"A{CH0+20}")
c4 = LineChart(); c4.title = "Account – XMC attach % and E-SIM activation % (fields live since 24 Aug)"; c4.height = 9; c4.width = 16; c4.style = 10
c4.add_data(Reference(da, min_col=15, min_row=3, max_row=DAL), titles_from_data=True); c4.add_data(Reference(da, min_col=18, min_row=3, max_row=DAL), titles_from_data=True)
c4.set_categories(Reference(da, min_col=1, min_row=DA1, max_row=DAL)); c4.y_axis.number_format = "0%"; c4.x_axis.number_format = "dd-mmm"; c4.legend.position = "b"
db.add_chart(c4, f"I{CH0+20}")
widths(db, {"A":11,"B":36,"C":26,"D":10,"E":12,"F":11,"G":11,"H":12,"I":36,"J":17,"K":13,"L":13,"M":12,"N":12})
db.freeze_panes = "A4"; db.sheet_properties.tabColor = NAVY

# ================================================================= Insights (python computed, static) + live formulas
ins = wb.create_sheet("Insights", 1)
ins.sheet_view.showGridLines = False
title(ins, "Insights & Recommendations", f"Static insights computed from the full tracker history (as of {MAXD:%d %b %Y}); the live block at the top follows the Dashboard selection.", 12)
ins.column_dimensions["A"].width = 4; ins.column_dimensions["B"].width = 150
ROW = [4]
def line(text, kind="body"):
    rr = ROW[0]; ROW[0] += 1
    ins.merge_cells(f"B{rr}:L{rr}") if False else None
    if kind == "head":
        put(ins, f"A{rr}", "", fl=TEAL, border=False); put(ins, f"B{rr}", text, F(12, True, "FFFFFF"), TEAL); ins.row_dimensions[rr].height = 22
    elif kind == "live":
        put(ins, f"A{rr}", "●", F(10, True, TEAL), border=False); c = put(ins, f"B{rr}", text, F(11, True, NAVY), GREY); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    else:
        put(ins, f"A{rr}", "•", F(10, True, NAVY), border=False); c = put(ins, f"B{rr}", text, F(11), border=False); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ins.row_dimensions[rr].height = 32
def gap(): ROW[0] += 1

# computations
v = s[s.dup == 0]
daily = v.groupby(v["Date"].dt.date)["units"].sum().reindex(dates, fill_value=0)
mon = v.groupby(v["Date"].dt.to_period("M"))["units"].sum(); mdays = v.groupby(v["Date"].dt.to_period("M"))["Date"].nunique(); mavg = (mon / mdays)
last7 = daily.iloc[-7:].sum(); prev7 = daily.iloc[-14:-7].sum(); l7c = (last7 - prev7) / prev7
last30 = v[v["Date"] >= pd.Timestamp(MAXD) - pd.Timedelta(days=29)]; prev30 = v[(v["Date"] < pd.Timestamp(MAXD) - pd.Timedelta(days=29)) & (v["Date"] >= pd.Timestamp(MAXD) - pd.Timedelta(days=59))]
peak = daily.idxmax(); wk = v.groupby(v["Date"].dt.day_name())["units"].sum() / v.groupby(v["Date"].dt.day_name())["Date"].nunique()
gshare = v["gsr"].sum() / v["units"].sum()
gt = {"Speed Upgrade": v.t_speed.sum(), "X-PRO": v.t_xpro.sum(), "Adding LOB (Internet-TV)": v.t_lob.sum(), "Mobile Plus Plan": v.t_mplus.sum(), "XMC": v.t_xmc.sum()}
xmc = v[v["Did you add XMC?"].isin(["Yes","No"])]; xmc_rate = (xmc["Did you add XMC?"] == "Yes").mean()
xmc30 = last30[last30["Did you add XMC?"].isin(["Yes","No"])]; xmc30_rate = (xmc30["Did you add XMC?"] == "Yes").mean()
es = v[v["Did you Activate the E-SIM?"].isin(["Yes","No"])]; es_rate = (es["Did you Activate the E-SIM?"] == "Yes").mean()
es30 = last30[last30["Did you Activate the E-SIM?"].isin(["Yes","No"])]; es30_rate = (es30["Did you Activate the E-SIM?"] == "Yes").mean()
dev = v[v.device.isin(["iPhone","Samsung","Apple Watch","Samsung Watch","iPad","Motorola","BYOD"])]["device"].value_counts()
hc = r.groupby("Team Leader").size()
tl30 = last30[last30.in_roster == 1].groupby("TL")["units"].sum().reindex(tls, fill_value=0); tl30p = tl30 / hc
tlprev = prev30[prev30.in_roster == 1].groupby("TL")["units"].sum().reindex(tls, fill_value=0)
part30 = last30[last30.in_roster == 1].groupby("TL")["hr"].nunique().reindex(tls, fill_value=0) / hc
om30 = last30[last30.in_roster == 1].groupby("OM")["units"].sum(); omhc = r.groupby("Ops Manager").size(); om30p = (om30 / omhc)
ag30 = last30[last30.in_roster == 1].groupby("hr")["units"].sum().sort_values(ascending=False)
top10share = ag30.head(10).sum() / last30["units"].sum()
prod = r[r.Phase == "Production"]; zero_prod = prod[~prod.HRID.isin(set(ag30[ag30 > 0].index))]
train = r[r.Phase == "Training"]; train_units = last30[last30.hr.isin(set(train.HRID))]["units"].sum()
unm = s[s.in_roster == 0]; unm_share = len(unm) / len(s); unm_units = unm.units.sum()
junk = (s.order_valid == 0).sum()
hours = s["ct"].dt.hour.value_counts(); 
top_tl = tl30p.sort_values(ascending=False)
best_prod_tl = tl30p[[t for t in tls if (r[(r["Team Leader"]==t)&(r.Phase=="Production")].shape[0] > 0)]]
ten = last30[last30.in_roster == 1].groupby("Tenure")["units"].sum(); tenhc = r.groupby("Tenure Bucket").size(); tenp = (ten / tenhc).dropna()

line("LIVE – follows the Dashboard selection (dates / Ops Manager / Team Leader)", "head")
line(f'="Selected period "&TEXT({D_FROM},"dd-mmm")&" to "&TEXT({D_TO},"dd-mmm")&": "&TEXT(Dashboard!E5,"#,##0")&" units ("&TEXT(Dashboard!E6,"+#,##0;-#,##0;0")&" vs prior period), "&TEXT(Dashboard!M5,"#,##0")&" active sellers out of "&{HC_SC}&" agents in scope ("&TEXT(Dashboard!G10,"0%")&" participation)."', "live")
line(f'="Best leader in the period: "&INDEX(Dashboard!$B${LS1}:$B${LSL},MATCH(MAX(Dashboard!$H${LS1}:$H${LSL}),Dashboard!$H${LS1}:$H${LSL},0))&" with "&TEXT(MAX(Dashboard!$H${LS1}:$H${LSL}),"#,##0")&" units ("&TEXT(INDEX(Dashboard!$I${LS1}:$I${LSL},MATCH(MAX(Dashboard!$H${LS1}:$H${LSL}),Dashboard!$H${LS1}:$H${LSL},0)),"0.00")&" per head). Lowest: "&INDEX(Dashboard!$B${LS1}:$B${LSL},MATCH(MIN(Dashboard!$H${LS1}:$H${LSL}),Dashboard!$H${LS1}:$H${LSL},0))&" with "&TEXT(MIN(Dashboard!$H${LS1}:$H${LSL}),"#,##0")&" units."', "live")
line(f'="Top agent in scope: "&Dashboard!B{TB0+2}&" ("&Dashboard!C{TB0+2}&") with "&TEXT(Dashboard!D{TB0+2},"#,##0")&" units. Production agents in scope with zero units: "&COUNTIFS({asr("T")},1,{asr("E")},"Production",{asr("K")},0)&"."', "live")
line(f'="Mix in period: "&TEXT(IFERROR(Dashboard!I5/Dashboard!E5,0),"0%")&" GSR / "&TEXT(IFERROR(Dashboard!G5/Dashboard!E5,0),"0%")&" XM Lines. XMC attach "&TEXT(Dashboard!I10,"0.0%")&" ("&TEXT(Dashboard!I11,"+0.0%;-0.0%")&" vs prior), E-SIM activation "&TEXT(Dashboard!K10,"0.0%")&"."', "live")
gap()
line("1. ACCOUNT TREND (day-over-day)", "head")
line(f"Volume has stepped up every month: average units per day went {' → '.join(f'{m.strftime(chr(37)+chr(98))} {x:.1f}' for m, x in mavg.items())}. The tracker started on {min(dates):%d %b}, so July is a partial month.")
line(f"Last 7 days ({dates[-7]:%d %b}–{MAXD:%d %b}): {int(last7):,} units vs {int(prev7):,} the 7 days before ({l7c:+.0%}). Peak day so far: {peak:%a %d %b} with {int(daily.max()):,} units; {MAXD:%d %b} closed at {int(daily.iloc[-1]):,}.")
line(f"Weekday pattern (avg units per day): strongest {wk.idxmax()} ({wk.max():.0f}), weakest {wk.idxmin()} ({wk.min():.0f}). Use the Daily_Account WoW % column to compare like-for-like weekdays instead of raw DoD on weekends.")
line(f"Entries are logged between 14:00 and 07:00 Cairo time (US shift). A 'day' in this file is the form completion date, so late-shift sales after midnight count toward the next calendar day – keep that in mind when reading DoD swings.")
gap()
line("2. PRODUCT MIX & QUALITY OF SALE", "head")
line(f"GSRs are {gshare:.0%} of all units ({int(v.gsr.sum()):,}) vs XM Lines {1-gshare:.0%} ({int(v.xm.sum()):,}). GSR type mix: " + ", ".join(f"{k} {int(x):,}" for k, x in sorted(gt.items(), key=lambda z: -z[1])) + ". Speed Upgrade alone is {:.0%} of GSR type counts.".format(gt['Speed Upgrade']/sum(gt.values())))
line(f"XMC attach on XM Line sales is only {xmc_rate:.1%} overall ({int((xmc['Did you add XMC?']=='Yes').sum())} of {len(xmc):,} answered) and {xmc30_rate:.1%} in the last 30 days – the single biggest upsell gap. E-SIM activation on e-SIM sales is {es_rate:.1%} overall and {es30_rate:.1%} in the last 30 days.")
line(f"Device attach: " + ", ".join(f"{k} {int(x):,}" for k, x in dev.items()) + f". 'Free Line for 1 Year' is the dominant promotion ({int((v.promo=='Free Line for 1 Year').sum()):,} entries).")
gap()
line("3. LEADER LEVEL (last 30 days, units per head so team sizes are comparable)", "head")
line("Top 3 leaders by units per head: " + ", ".join(f"{t} {top_tl[t]:.1f} (HC {hc[t]}, participation {part30[t]:.0%})" for t in top_tl.index[:3]) + ".")
line("Bottom 3 leaders by units per head: " + ", ".join(f"{t} {top_tl[t]:.1f} (HC {hc[t]}, participation {part30[t]:.0%})" for t in top_tl.index[-3:]) + ". Note: teams with Training-phase agents will naturally sit lower – check the Phase column in Agent_Summary before comparing.")
mv = (tl30 - tlprev); 
line("Biggest 30-day movers vs the prior 30 days – up: " + ", ".join(f"{t} {int(mv[t]):+,}" for t in mv.sort_values(ascending=False).index[:3] if mv[t] > 0) + ". Down: " + (", ".join(f"{t} {int(mv[t]):+,}" for t in mv.sort_values().index[:3] if mv[t] < 0) or "none") + ".")
line("Ops Managers (units per head, last 30 days): " + ", ".join(f"{o} {x:.1f} (HC {omhc[o]})" for o, x in om30p.sort_values(ascending=False).items()) + ".")
line("Tenure effect (units per head, last 30 days): " + ", ".join(f"{k}: {x:.1f}" for k, x in tenp.sort_values(ascending=False).items()) + ". Newer agents (30-60 days) are the largest group on the roster – ramp coaching has the biggest leverage.")
gap()
line("4. AGENT LEVEL", "head")
line(f"{'Concentration risk:' if top10share > 0.35 else 'Output is well spread:'} the top 10 agents produced {top10share:.0%} of all units in the last 30 days. #1 is {rmap.loc[ag30.index[0],'Agent Name']} ({rmap.loc[ag30.index[0],'Team Leader']}) with {int(ag30.iloc[0]):,} units – {ag30.iloc[0]/ag30.sum():.0%} of the roster total alone.")
line(f"{len(zero_prod)} of {len(prod)} Production-phase agents logged zero units in the last 30 days (" + ", ".join(f"{t}: {n}" for t, n in zero_prod.groupby('Team Leader').size().sort_values(ascending=False).head(5).items()) + " …). These are the first names to pull into the Bottom-10 coaching list on the Dashboard.")
line(f"Training-phase agents ({len(train)} on roster) already contributed {int(train_units):,} units in the last 30 days – track them separately so early wins are recognised and tracker discipline is built from day one.")
gap()
line("5. DATA QUALITY (fix these to make the numbers trustworthy)", "head")
line(f"{len(unm):,} entries ({unm_share:.0%}) carry an HRID that is not on the roster ({unm.hr_raw.nunique()} distinct values, {int(unm_units):,} units). Most are 9-digit IDs that look valid (likely attrited or transferred agents) – e.g. HRID {unm.hr_raw.value_counts().index[0]} alone has {unm.hr_raw.value_counts().iloc[0]} entries. Add these agents to the roster (or a 'Leavers' tab) to attribute them to leaders.")
line(f"{int(s.dup.sum()):,} entries are exact duplicates (same HRID + account + order number + type) and are excluded by default (toggle on the Dashboard). {junk:,} entries have an unusable order number (e.g. '1', '0', '.', 'N/A', pasted work-order text). Make the Order Number field numeric-only in the Form.")
line(f"The 'Count of GSR' field only exists since 24 Aug; before that GSR units are inferred from the number of GSR types ticked. 8-digit / 10-digit HRID typos were auto-corrected where a unique roster match existed; make HRID a validated 9-digit field in the Form.")
gap()
line("6. RECOMMENDED ACTIONS FOR THE 3 NEWLY-PROMOTED LEADERS", "head")
line("Daily (10 min): open Daily_Agent with your team selected, read the DoD Δ row and the Rank-in-Team row; message every agent with 0 units by mid-shift and recognise the top mover publicly.")
line("Weekly: use the Dashboard with From/To set to the week and Team Leader set to your name – review Participation %, Units / Active Seller and the Bottom-10 list; pair each bottom-10 agent with a top-10 agent for a side-by-side on XMC attach and e-SIM activation.")
line("Pitch focus: XMC attach is below 15% – build a one-line XMC offer into every XM Line close; e-SIM activation misses should be followed up the same shift, since an unactivated e-SIM is a sale at risk.")
line("Hygiene: every agent submits the form before closing the interaction, HRID copied from the roster, order number from the system – one wrong HRID removes the sale from your scorecard.")
ins.sheet_properties.tabColor = "548235"

# ================================================================= Guide
gd = wb.create_sheet("Guide")
gd.sheet_view.showGridLines = False
title(gd, "How to use this workbook", None, 3)
gd.column_dimensions["A"].width = 24; gd.column_dimensions["B"].width = 120
rows_g = [
 ("Dashboard", "Control panel in C5:C9 (yellow). From/To dates drive every period figure; Ops Manager / Team Leader narrow the KPI tiles, Agent Summary, Top/Bottom lists and Live Insights. 'Exclude Duplicates' removes repeated submissions of the same order."),
 ("Daily_Account", "Whole-account day-over-day for every day in the tracker: units, DoD Δ and %, trailing 7-day average, week-over-week (same weekday), active sellers, XMC attach, e-SIM activation, GSR type split, devices and data-quality counts."),
 ("Daily_Leader", "Metric dropdown (B3). Summary rows 6-12 use the Dashboard date range; the matrix below shows each day for every Team Leader, with a day-over-day change block on the right."),
 ("Daily_Agent", "Team Leader dropdown (B3) + metric (B4). Agents of the chosen team populate automatically with period totals, last-day vs prior-day, rank in team, active days, then a daily matrix and a DoD block."),
 ("Agent_Trend", "One agent (B3): daily submissions / XM / GSR / units, DoD, 7-day average, cumulative and a comparison against the team's average per agent, with a chart."),
 ("Agent_Summary", "All roster agents for the Dashboard period & scope. Filter/sort with the header arrows. 'Rank (Scope)' ranks inside the current filter; 'Rank (Team)' ranks inside the agent's team. Red shading = no sale for more than 7 days."),
 ("Insights", "Top block is live (formulas). Numbered sections are the analyst commentary computed from the full history at build time."),
 ("Sales_Data", "Cleaned tracker: one row per form submission with mapped HRID, leader, ops manager, derived units and flags. 'Include' is the only formula column (duplicate toggle). Add new rows here to extend the data – keep the same columns and extend the formula ranges if you exceed row " + str(LAST) + "."),
 ("Roster / Lists", "Roster drives agent ↔ leader mapping and the dropdown lists. Team Key is a helper for the Daily_Agent sheet."),
 ("", ""),
 ("Definitions", ""),
 ("Units", "XM Lines sold + GSR units. XM Lines = 'No. of XM Lines' (1 if blank). GSR units = 'Count of GSR' when present, otherwise the number of GSR types ticked (field only exists since 24 Aug 2026)."),
 ("Submissions", "Number of form entries (after the duplicate toggle)."),
 ("Active Sellers / Participation", "Roster agents with at least one unit in the period; participation = active ÷ headcount in scope. Agents not on the roster are shown separately as 'Not in Roster'."),
 ("XMC Attach %", "'Did you add XMC?' = Yes ÷ (Yes + No). E-SIM Activation % = 'Did you Activate the E-SIM?' Yes ÷ (Yes + No)."),
 ("Prior period", "The same number of days immediately before the From Date (shown in C11)."),
 ("Duplicate", "Same mapped HRID + account number + order number + entry type, where the order number is a 6-20 digit number. The first occurrence is kept."),
 ("HRID mapping", "Digits only; 8-digit IDs starting with 100 and 10-digit IDs starting with 1000 are repaired when a unique roster match exists. Everything else is 'Not in Roster'."),
 ("Fonts & layout", "Calibri throughout, all cells centred, yellow = input cells, navy headers, green/red = day-over-day up/down."),
]
for i, (a, b) in enumerate(rows_g, start=3):
    if a and not b: put(gd, f"A{i}", a, F(12, True, "FFFFFF"), TEAL); put(gd, f"B{i}", "", fl=TEAL)
    elif a: put(gd, f"A{i}", a, F(11, True), LIGHT); c = put(gd, f"B{i}", b, F(10)); c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); gd.row_dimensions[i].height = 34
gd.sheet_properties.tabColor = "548235"

# ================================================================= global: Calibri + centre everything
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f0 = c.font
            c.font = Font(name="Calibri", size=f0.size or 11, bold=f0.bold, italic=f0.italic, color=f0.color)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=(c.alignment.wrap_text or False))
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
